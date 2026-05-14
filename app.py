import os
import io
import re
import json
import tempfile
from pathlib import Path
from flask import Flask, render_template, request, jsonify, send_file
import anthropic

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50MB max


# ──────────────────────────────────────────────
#  TEXT EXTRACTION
# ──────────────────────────────────────────────

def extract_text_from_pdf(path):
    import pdfplumber
    text = ""
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            t = page.extract_text()
            if t:
                text += f"\n[Página {i}]\n{t}\n"
    return text


def extract_text_from_docx(path):
    from docx import Document
    doc = Document(path)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def extract_text_from_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    text = ""
    for i, slide in enumerate(prs.slides, 1):
        text += f"\n[Diapositiva {i}]\n"
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text += shape.text + "\n"
    return text


def extract_text(path, filename):
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(path)
    if ext in (".docx", ".doc"):
        return extract_text_from_docx(path)
    if ext in (".pptx", ".ppt"):
        return extract_text_from_pptx(path)
    if ext == ".txt":
        return Path(path).read_text(encoding="utf-8", errors="ignore")
    return ""


# ──────────────────────────────────────────────
#  FILE GENERATION
# ──────────────────────────────────────────────

def clean_markdown(text):
    """Strip markdown symbols for plain-text writers."""
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    text = re.sub(r'`(.*?)`', r'\1', text)
    return text


def generate_pdf(markdown_text, title="Apuntes IA"):
    from fpdf import FPDF

    class PDF(FPDF):
        def header(self):
            self.set_font("Helvetica", "B", 10)
            self.set_text_color(120, 80, 220)
            self.cell(0, 8, title, align="C")
            self.set_draw_color(120, 80, 220)
            self.set_line_width(0.5)
            self.line(10, 16, 200, 16)
            self.ln(6)

        def footer(self):
            self.set_y(-13)
            self.set_font("Helvetica", "I", 8)
            self.set_text_color(150, 150, 150)
            self.cell(0, 8, f"Generado por ApuntesIA  •  Página {self.page_no()}", align="C")

    pdf = PDF()
    pdf.set_margins(15, 20, 15)
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=18)

    def safe(s):
        return s.encode("latin-1", errors="replace").decode("latin-1")

    for line in markdown_text.split("\n"):
        stripped = line.rstrip()
        if stripped.startswith("# "):
            pdf.set_font("Helvetica", "B", 17)
            pdf.set_text_color(80, 40, 180)
            pdf.multi_cell(0, 10, safe(stripped[2:]))
            pdf.set_text_color(0, 0, 0)
            pdf.ln(1)
        elif stripped.startswith("## "):
            pdf.set_font("Helvetica", "B", 14)
            pdf.set_text_color(50, 100, 200)
            pdf.multi_cell(0, 8, safe(stripped[3:]))
            pdf.set_text_color(0, 0, 0)
            pdf.ln(1)
        elif stripped.startswith("### "):
            pdf.set_font("Helvetica", "B", 12)
            pdf.set_text_color(30, 140, 160)
            pdf.multi_cell(0, 7, safe(stripped[4:]))
            pdf.set_text_color(0, 0, 0)
        elif stripped.startswith("#### "):
            pdf.set_font("Helvetica", "BI", 11)
            pdf.set_text_color(60, 60, 60)
            pdf.multi_cell(0, 6, safe(stripped[5:]))
            pdf.set_text_color(0, 0, 0)
        elif stripped.startswith(("- ", "* ", "• ")):
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(0, 5.5, safe("  •  " + clean_markdown(stripped[2:])))
        elif re.match(r"^\d+\. ", stripped):
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(0, 5.5, safe("     " + clean_markdown(stripped)))
        elif stripped == "" or stripped == "---":
            pdf.ln(3)
        else:
            # inline bold
            parts = re.split(r"\*\*(.*?)\*\*", stripped)
            pdf.set_font("Helvetica", "", 10)
            for i, part in enumerate(parts):
                if not part:
                    continue
                if i % 2 == 1:
                    pdf.set_font("Helvetica", "B", 10)
                else:
                    pdf.set_font("Helvetica", "", 10)
                # write inline — can't do true inline with FPDF easily, fallback
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(0, 5.5, safe(clean_markdown(stripped)))

    buf = io.BytesIO()
    pdf.output(buf)
    buf.seek(0)
    return buf, "application/pdf", "apuntes.pdf"


def generate_docx(markdown_text, title="Apuntes IA"):
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    # Title
    t = doc.add_heading(title, level=0)
    t.runs[0].font.color.rgb = RGBColor(80, 40, 180)

    for line in markdown_text.split("\n"):
        s = line.rstrip()
        if s.startswith("# "):
            h = doc.add_heading(s[2:], level=1)
            if h.runs:
                h.runs[0].font.color.rgb = RGBColor(80, 40, 180)
        elif s.startswith("## "):
            h = doc.add_heading(s[3:], level=2)
            if h.runs:
                h.runs[0].font.color.rgb = RGBColor(50, 100, 200)
        elif s.startswith("### "):
            doc.add_heading(s[4:], level=3)
        elif s.startswith("#### "):
            doc.add_heading(s[5:], level=4)
        elif s.startswith(("- ", "* ", "• ")):
            doc.add_paragraph(clean_markdown(s[2:]), style="List Bullet")
        elif re.match(r"^\d+\. ", s):
            doc.add_paragraph(clean_markdown(re.sub(r"^\d+\. ", "", s)), style="List Number")
        elif s.strip() == "" or s.strip() == "---":
            doc.add_paragraph("")
        else:
            p = doc.add_paragraph()
            parts = re.split(r"\*\*(.*?)\*\*", s)
            for i, part in enumerate(parts):
                if not part:
                    continue
                run = p.add_run(part)
                run.font.size = Pt(11)
                if i % 2 == 1:
                    run.bold = True

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "apuntes.docx"


def generate_pptx(markdown_text, title="Apuntes IA"):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    PURPLE = RGBColor(80, 40, 180)
    BLUE = RGBColor(50, 100, 200)
    WHITE = RGBColor(255, 255, 255)
    BG = RGBColor(15, 10, 35)

    def set_bg(slide, color):
        from pptx.oxml.ns import qn
        from lxml import etree
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = color

    # ── TITLE SLIDE ──
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)
    set_bg(slide, BG)

    txb = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.33), Inches(1.5))
    tf = txb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = title
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.color.rgb = WHITE

    txb2 = slide.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11.33), Inches(0.8))
    tf2 = txb2.text_frame
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "Generado con ApuntesIA  •  Claude AI"
    r2.font.size = Pt(18)
    r2.font.color.rgb = RGBColor(160, 130, 240)

    # ── PARSE SECTIONS ──
    sections = []
    current_title = "Introducción"
    current_bullets = []

    for line in markdown_text.split("\n"):
        s = line.rstrip()
        if s.startswith("# ") or s.startswith("## "):
            if current_bullets:
                sections.append((current_title, current_bullets[:]))
            current_title = s.lstrip("# ").strip()
            current_bullets = []
        elif s.startswith("### ") or s.startswith("#### "):
            current_bullets.append(("sub", s.lstrip("# ").strip()))
        elif s.startswith(("- ", "* ", "• ")):
            current_bullets.append(("bullet", clean_markdown(s[2:])))
        elif re.match(r"^\d+\. ", s):
            current_bullets.append(("bullet", clean_markdown(re.sub(r"^\d+\. ", "", s))))
        elif s.strip() and s.strip() != "---":
            current_bullets.append(("text", clean_markdown(s)))

    if current_bullets:
        sections.append((current_title, current_bullets))

    # ── CONTENT SLIDES ──
    for sec_title, bullets in sections[:25]:
        layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(layout)
        set_bg(slide, BG)

        # Title bar
        title_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(12.53), Inches(1))
        ttf = title_box.text_frame
        tp = ttf.paragraphs[0]
        tr = tp.add_run()
        tr.text = sec_title[:90]
        tr.font.size = Pt(28)
        tr.font.bold = True
        tr.font.color.rgb = RGBColor(160, 130, 240)

        # Divider line approximated with a colored rectangle
        line_box = slide.shapes.add_textbox(Inches(0.4), Inches(1.3), Inches(12.53), Inches(0.05))
        line_box.fill.solid()
        line_box.fill.fore_color.rgb = PURPLE

        # Content
        content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12.33), Inches(5.7))
        ctf = content_box.text_frame
        ctf.word_wrap = True

        shown = 0
        for i, (kind, btext) in enumerate(bullets[:16]):
            p = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
            run = p.add_run()
            if kind == "sub":
                run.text = "▸  " + btext[:120]
                run.font.size = Pt(15)
                run.font.bold = True
                run.font.color.rgb = RGBColor(120, 200, 240)
            elif kind == "bullet":
                run.text = "•  " + btext[:130]
                run.font.size = Pt(13)
                run.font.color.rgb = WHITE
            else:
                run.text = btext[:140]
                run.font.size = Pt(12)
                run.font.color.rgb = RGBColor(200, 200, 220)
            shown += 1

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf, "application/vnd.openxmlformats-officedocument.presentationml.presentation", "apuntes.pptx"


def build_output(fmt, markdown_text, title="Apuntes IA"):
    if fmt == "pdf":
        return generate_pdf(markdown_text, title)
    if fmt == "docx":
        return generate_docx(markdown_text, title)
    if fmt == "pptx":
        return generate_pptx(markdown_text, title)
    buf = io.BytesIO(markdown_text.encode())
    buf.seek(0)
    return buf, "text/plain", "apuntes.txt"


# ──────────────────────────────────────────────
#  ROUTES
# ──────────────────────────────────────────────

@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/process", methods=["POST"])
def process_document():
    api_key = request.form.get("api_key", "").strip()
    instructions = request.form.get("instructions", "").strip()
    output_format = request.form.get("output_format", "pdf")
    detail_level = request.form.get("detail_level", "ultra")

    if not api_key:
        return jsonify({"error": "Se requiere una API Key de Anthropic."}), 400

    files = request.files.getlist("files")
    if not files or all(f.filename == "" for f in files):
        return jsonify({"error": "No se subieron archivos."}), 400

    all_text = ""
    names = []
    for f in files:
        if not f.filename:
            continue
        suffix = Path(f.filename).suffix
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            f.save(tmp.name)
            t = extract_text(tmp.name, f.filename)
            all_text += f"\n\n{'='*60}\nARCHIVO: {f.filename}\n{'='*60}\n{t}"
            names.append(f.filename)
            os.unlink(tmp.name)

    if not all_text.strip():
        return jsonify({"error": "No se pudo extraer texto de los archivos."}), 400

    detail_instructions = {
        "ultra": "Máximo detalle posible. Cada concepto explicado exhaustivamente con ejemplos, contexto histórico si aplica, implicaciones prácticas y conexiones con otros temas.",
        "high": "Alto nivel de detalle. Explica todos los conceptos importantes con ejemplos claros.",
        "medium": "Detalle moderado. Cubre los puntos principales con explicaciones concisas.",
        "summary": "Resumen ejecutivo. Solo los puntos más críticos y esenciales.",
    }.get(detail_level, "Máximo detalle.")

    system = f"""Eres el mejor creador de apuntes académicos del mundo. Tus apuntes son legendarios por su claridad, profundidad y utilidad pedagógica.

REGLAS ABSOLUTAS:
1. Usa Markdown rico: # ## ### ####, **negrita**, listas anidadas, tablas, separadores ---
2. Nivel de detalle: {detail_instructions}
3. Estructura OBLIGATORIA:
   - 📚 Título principal
   - 📋 Resumen ejecutivo (3-5 oraciones)
   - 🎯 Objetivos de aprendizaje (lista)
   - 📖 Contenido (múltiples secciones temáticas detalladas)
   - 💡 Conceptos clave (definiciones precisas)
   - 📝 Ejemplos y casos prácticos
   - ⚠️ Puntos críticos a recordar
   - 🔗 Conexiones y conclusiones
4. Dentro de cada sección: usa subsecciones ###, listas con guiones, tablas comparativas cuando proceda.
5. Añade contexto, ejemplos del mundo real, analogías y tips mnemotécnicos.
6. Usa emojis estratégicos para mejorar la escaneabilidad visual.
7. Al final incluye un glosario con los términos técnicos."""

    user_msg = f"""Crea apuntes ultra detallados del siguiente material.

INSTRUCCIONES ESPECÍFICAS:
{instructions if instructions else 'Genera apuntes completos y exhaustivos de todo el contenido proporcionado.'}

ARCHIVOS PROCESADOS: {', '.join(names)}

CONTENIDO:
{all_text[:60000]}"""

    try:
        client = anthropic.Anthropic(api_key=api_key)
        response = client.messages.create(
            model="claude-opus-4-7",
            max_tokens=8000,
            system=system,
            messages=[{"role": "user", "content": user_msg}],
        )
        notes = response.content[0].text
    except anthropic.AuthenticationError:
        return jsonify({"error": "API Key inválida. Verifica tu clave de Anthropic."}), 401
    except anthropic.RateLimitError:
        return jsonify({"error": "Límite de uso alcanzado. Intenta más tarde."}), 429
    except Exception as e:
        return jsonify({"error": f"Error al llamar a Claude: {str(e)}"}), 500

    return jsonify({"success": True, "notes": notes, "output_format": output_format})


@app.route("/api/download", methods=["POST"])
def download():
    data = request.get_json()
    notes = data.get("notes", "")
    fmt = data.get("output_format", "pdf")
    title = data.get("title", "Apuntes IA")

    buf, mimetype, filename = build_output(fmt, notes, title)
    return send_file(buf, mimetype=mimetype, as_attachment=True, download_name=filename)


@app.route("/api/text-tool", methods=["POST"])
def text_tool():
    data = request.get_json()
    api_key = data.get("api_key", "").strip()
    tool = data.get("tool", "")
    text = data.get("text", "")
    text2 = data.get("text2", "")
    options = data.get("options", {})
    output_format = data.get("output_format", "txt")

    if not api_key:
        return jsonify({"error": "Se requiere API Key."}), 400

    prompts = {
        "compress": f"Comprime y resume el siguiente texto al máximo, manteniendo TODA la información esencial. Usa bullet points y sé muy conciso:\n\n{text}",
        "expand": f"Expande y desarrolla el siguiente texto en profundidad. Añade contexto, ejemplos, explicaciones detalladas y secciones bien estructuradas:\n\n{text}",
        "extract": f"Extrae y lista TODOS los puntos clave, conceptos importantes y datos relevantes del siguiente texto. Organízalos por categorías:\n\n{text}",
        "merge": f"Combina y unifica coherentemente los siguientes dos textos en uno solo, bien estructurado y sin repeticiones:\n\nTEXTO 1:\n{text}\n\nTEXTO 2:\n{text2}",
        "clean": f"Limpia, corrige, formatea y mejora la redacción del siguiente texto. Corrige errores, mejora el estilo y la estructura:\n\n{text}",
        "split": f"Divide el siguiente texto en secciones temáticas bien definidas con títulos claros y numerados:\n\n{text}",
        "rewrite": f"Reescribe el siguiente texto con estilo {options.get('style', 'académico y profesional')}, mejorando la claridad y el impacto:\n\n{text}",
        "translate": f"Traduce el siguiente texto al {options.get('language', 'inglés')} de forma natural y precisa:\n\n{text}",
        "bullets": f"Convierte el siguiente texto en una lista de bullet points ultra clara y estructurada, organizados por temas:\n\n{text}",
        "table": f"Convierte la información del siguiente texto en una o más tablas markdown bien organizadas:\n\n{text}",
    }

    prompt = prompts.get(tool, f"Procesa el siguiente texto:\n\n{text}")

    try:
        client = anthropic.Anthropic(api_key=api_key)
        response = client.messages.create(
            model="claude-opus-4-7",
            max_tokens=4000,
            messages=[{"role": "user", "content": prompt}],
        )
        result = response.content[0].text
    except anthropic.AuthenticationError:
        return jsonify({"error": "API Key inválida."}), 401
    except Exception as e:
        return jsonify({"error": str(e)}), 500

    return jsonify({"success": True, "result": result, "output_format": output_format})


@app.route("/api/download-text", methods=["POST"])
def download_text():
    data = request.get_json()
    text = data.get("text", "")
    fmt = data.get("output_format", "txt")
    title = data.get("title", "Resultado")

    buf, mimetype, filename = build_output(fmt, text, title)
    return send_file(buf, mimetype=mimetype, as_attachment=True, download_name=filename)


if __name__ == "__main__":
    app.run(debug=True, port=5000)
